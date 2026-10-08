import streamlit as st
from streamlit_drawable_canvas import st_canvas
import pandas as pd
from PIL import Image
import io
import base64
import hashlib
from datetime import datetime, date
import pytz
import numpy as np
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as OpenPyxlImage
import time
import uuid
import json
import gspread
from google.oauth2.service_account import Credentials
from fpdf import FPDF
import tempfile
import os

# --- CONFIGURAÇÃO DA PÁGINA ---
st.set_page_config(
    page_title="Habituação de Atiradores",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# --- CONFIGURAÇÃO DE FUSO HORÁRIO BRASÍLIA/SÃO PAULO ---
FUSO_SP = pytz.timezone("America/Sao_Paulo")

def obter_data_hora_atual():
    return datetime.now(FUSO_SP)

# --- CSS PARA FORÇAR CAIXA ALTA (MAIÚSCULAS) EM TEMPO REAL NOS CAMPOS DE TEXTO ---
st.markdown(
    """
    <style>
    input[type="text"] {
        text-transform: uppercase;
    }
    </style>
    """,
    unsafe_allow_html=True
)

# --- DADOS INSTITUCIONAIS DA ENTIDADE DE TIRO (PARA RELATÓRIO OFICIAL SFPC) ---
NOME_ENTIDADE_TIRO = "CLUBE DE CAÇA E TIRO URBANO"
CR_ENTIDADE_TIRO = "123456"
CNPJ_ENTIDADE_TIRO = "00.000.000/0001-00"
ENDERECO_ENTIDADE_TIRO = "RUA DO CLUBE, Nº 100 - CENTRO"
SFPC_VINCULACAO = "5ª RM / SFPC"
CIDADE_UF_ENTIDADE = "ITAPEJARA D'OESTE - PR"
NOME_RESPONSAVEL_ENTIDADE = "RODRIGO HENRIQUE NEVES"

# --- NOME DO FICHEIRO DA LOGO ---
PATH_LOGO = "LOGO_CCTU-removebg-preview.png"

def obter_logo_base64():
    path_final = None
    if os.path.exists(PATH_LOGO):
        path_final = PATH_LOGO
    elif os.path.exists("LOGO_CCTU-removebg-preview.jpg"):
        path_final = "LOGO_CCTU-removebg-preview.jpg"

    if path_final:
        try:
            with open(path_final, "rb") as image_file:
                return base64.b64encode(image_file.read()).decode('utf-8')
        except Exception:
            return ""
    return ""

LOGO_B64 = obter_logo_base64()

# --- AUTENTICAÇÃO E CONEXÃO COM GOOGLE SHEETS ---
SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive"
]

@st.cache_resource
def get_gspread_client():
    if "gcp_service_account" not in st.secrets:
        st.error("A seção [gcp_service_account] não foi encontrada nos Secrets do Streamlit.")
        st.stop()

    sec_data = st.secrets["gcp_service_account"]

    if "json_data" in sec_data:
        raw_json = sec_data["json_data"]
        if isinstance(raw_json, str):
            creds_info = json.loads(raw_json, strict=False)
        else:
            creds_info = dict(raw_json)
    else:
        creds_info = dict(sec_data)

    if "private_key" in creds_info:
        creds_info["private_key"] = creds_info["private_key"].replace("\\n", "\n")

    creds = Credentials.from_service_account_info(creds_info, scopes=SCOPES)
    client = gspread.authorize(creds)
    return client

def get_sheet():
    client = get_gspread_client()
    spreadsheet = client.open("Habituacoes_Clube")
    return spreadsheet.sheet1

# --- MAPEAMENTO DINÂMICO DE MUNIÇÕES POR TIPO DE ARMA ---
OPCOES_MUNICAO_POR_ARMA = {
    "Pistola": [
        "9mm Luger", ".380 ACP", ".40 S&W", ".45 ACP", ".38 TPC", 
        ".22 LR", ".38 Super Auto", "10mm Auto", "7,65mm Browning", "6,35mm (.25 ACP)"
    ],
    "Revólver": [
        ".38 SPL", ".357 Magnum", ".22 LR", ".32 S&W / Longo", 
        ".44 Magnum", ".454 Casull", ".22 WMR", ".44-40 WCF"
    ],
    "Carabina/Fuzil": [
        "5.56x45mm / .223 Rem", ".22 LR", "9mm Luger", ".300 Blackout", 
        ".308 Win / 7.62x51mm", ".38 SPL", ".357 Magnum", ".40 S&W", 
        ".380 ACP", ".17 HMR", "6.5 Creedmoor", ".44 Magnum", ".44-40 WCF", ".30-06 Springfield"
    ],
    "Espingarda": [
        "12 GA", "20 GA", "28 GA", "36 GA / .410", "16 GA", "24 GA", "32 GA"
    ]
}

def processar_imagem_canvas(img_array):
    try:
        img_uint8 = img_array.astype(np.uint8)
        img_pil = Image.fromarray(img_uint8, mode="RGBA")
        img_pil.thumbnail((300, 150))
        background = Image.new("RGB", img_pil.size, (255, 255, 255))
        background.paste(img_pil, mask=img_pil.split()[3])
        buffered = io.BytesIO()
        background.save(buffered, format="PNG", optimize=True)
        return base64.b64encode(buffered.getvalue()).decode('utf-8')
    except Exception as e:
        st.error(f"Erro ao processar imagem da assinatura: {e}")
        return ""

def salvar_habituação_sheets(nome, cr, sigma, arma, municao, qtd, img_array):
    assinatura_b64 = processar_imagem_canvas(img_array)
    if not assinatura_b64:
        raise ValueError("Não foi possível processar o desenho da assinatura.")

    data_hora_str = obter_data_hora_atual().strftime("%Y-%m-%d %H:%M:%S")
    registro_id = str(uuid.uuid4())[:8]

    nome = nome.strip().upper()
    cr = cr.strip().upper()
    sigma = sigma.strip().upper() if sigma else ""

    payload_validacao = f"{data_hora_str}|{nome}|{cr}|{sigma}|{arma}|{municao}|{qtd}|{assinatura_b64[:50]}"
    hash_integridade = hashlib.sha256(payload_validacao.encode('utf-8')).hexdigest()

    linha_dados = [
        registro_id,
        data_hora_str,
        nome,
        cr,
        sigma or "",
        arma,
        municao,
        int(qtd),
        assinatura_b64,
        hash_integridade
    ]

    sheet = get_sheet()
    sheet.append_row(linha_dados)

def carregar_habituacoes_sheets():
    try:
        sheet = get_sheet()
        registros = sheet.get_all_records()
        if not registros:
            return pd.DataFrame()
        return pd.DataFrame(registros)
    except Exception as e:
        st.error(f"Erro ao carregar dados do Google Sheets: {e}")
        return pd.DataFrame()

def obter_atiradores_existentes(df_hab):
    if df_hab.empty or 'nome_atirador' not in df_hab.columns:
        return []
    
    atiradores = []
    df_validos = df_hab.dropna(subset=['nome_atirador'])
    
    for nome, group in df_validos.groupby('nome_atirador'):
        nome_str = str(nome).strip().upper()
        if nome_str:
            cr_val = ""
            if 'cr_atirador' in group.columns and not group['cr_atirador'].dropna().empty:
                cr_val = str(group['cr_atirador'].dropna().iloc[-1]).strip().upper()
            
            sigma_val = ""
            if 'sigma_atirador' in group.columns and not group['sigma_atirador'].dropna().empty:
                sigma_val = str(group['sigma_atirador'].dropna().iloc[-1]).strip().upper()

            label = f"{nome_str} - CR: {cr_val}" if cr_val else nome_str
            atiradores.append({
                "nome": nome_str,
                "cr": cr_val,
                "sigma": sigma_val,
                "rotulo": label
            })
            
    return sorted(atiradores, key=lambda x: x['nome'])

def obter_valor_assinatura(row):
    colunas_possiveis = ['assinatura_base64a', 'assinatura_base64', 'assinatura', 'Assinatura Digital']
    for col in colunas_possiveis:
        if col in row and pd.notna(row[col]) and str(row[col]).strip() != "":
            val_str = str(row[col]).strip()
            if "base64," in val_str:
                try:
                    return val_str.split("base64,")[1].split('"')[0].split("'")[0].split(')')[0].strip()
                except Exception:
                    pass
            elif len(val_str) > 100 and not val_str.startswith("http") and not val_str.startswith("="):
                return val_str
    try:
        val_ind = row.iloc[8]
        if pd.notna(val_ind) and len(str(val_ind)) > 100:
            return str(val_ind).strip()
    except Exception:
        pass
    return ""

def gerar_excel_com_assinaturas(df_hab):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Habituações"
    headers = [
        "ID", "Data/Hora", "Nome Atirador", "CR", "SIGMA", 
        "Tipo Arma", "Tipo Munição", "Qtd. Munição", "Assinatura Digital", "Hash Integridade"
    ]
    ws.append(headers)
    ws.column_dimensions['I'].width = 25

    for row_idx, (_, row) in enumerate(df_hab.iterrows(), start=2):
        ws.cell(row=row_idx, column=1, value=row.get('id', ''))
        ws.cell(row=row_idx, column=2, value=row.get('data_hora', ''))
        ws.cell(row=row_idx, column=3, value=row.get('nome_atirador', ''))
        ws.cell(row=row_idx, column=4, value=row.get('cr_atirador', ''))
        ws.cell(row=row_idx, column=5, value=row.get('sigma_atirador', ''))
        ws.cell(row=row_idx, column=6, value=row.get('tipo_arma', ''))
        ws.cell(row=row_idx, column=7, value=row.get('tipo_municao', ''))
        ws.cell(row=row_idx, column=8, value=row.get('qtd_municao', ''))
        ws.cell(row=row_idx, column=10, value=row.get('hash_integridade', ''))
        ws.row_dimensions[row_idx].height = 55

        assinatura_b64 = obter_valor_assinatura(row)
        if assinatura_b64:
            try:
                img_bytes = base64.b64decode(assinatura_b64)
                img_file = io.BytesIO(img_bytes)
                img = OpenPyxlImage(img_file)
                img.width = 140
                img.height = 60
                ws.add_image(img, f"I{row_idx}")
            except Exception:
                pass

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def gerar_excel_modelo_sfpc(df_hab, periodo_mes_ano="MÊS DE ________ DE 2026"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relação de Atiradores"
    ws.views.sheetView[0].showGridLines = True

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    grey_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    
    font_bold_title = Font(name="Calibri", size=11, bold=True)
    font_header = Font(name="Calibri", size=10, bold=True)
    font_data = Font(name="Calibri", size=10)
    font_legal = Font(name="Calibri", size=10, bold=True)
    font_subtext = Font(name="Calibri", size=9, italic=True)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin_border_side = Side(border_style="thin", color="000000")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 45
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 28

    ws.merge_cells('A1:E1')
    cell_a1 = ws['A1']
    cell_a1.value = f"RELAÇÃO DOS ATIRADORES E DOS ATLETAS QUE FREQUENTARAM ESTA ENTIDADE DE TIRO DESPORTIVO NO {periodo_mes_ano}"
    cell_a1.font = font_bold_title
    cell_a1.fill = yellow_fill
    cell_a1.alignment = align_center

    ws['A2'] = "ORD\nNUMÉRICA"
    ws['B2'] = f"Nome da entidade de tiro: {NOME_ENTIDADE_TIRO}"
    ws['C2'] = f"CR da entidade Tiro: {CR_ENTIDADE_TIRO}\nCNPJ: {CNPJ_ENTIDADE_TIRO}"
    ws['D2'] = f"Endereço da entidade de tiro:\n{ENDERECO_ENTIDADE_TIRO}"
    ws['E2'] = f"SFPC de vinculação:\n{SFPC_VINCULACAO}"

    for col in ['A', 'B', 'C', 'D', 'E']:
        cell = ws[f'{col}2']
        cell.font = font_header
        cell.fill = yellow_fill
        cell.alignment = align_center

    ws['A3'] = ""
    ws['B3'] = "Nome Completo do atirador/atleta"
    ws['C3'] = "CR atirador/atleta"
    ws['D3'] = "CPF atirador/atleta"
    ws['E3'] = "Data em que frequentou a entidade"

    ws.merge_cells('A2:A3')

    for col in ['A', 'B', 'C', 'D', 'E']:
        cell = ws[f'{col}3']
        cell.font = font_header
        cell.fill = grey_fill
        cell.alignment = align_center

    for r in range(1, 4):
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = thin_border

    start_row = 4
    
    if not df_hab.empty:
        for i, (_, row) in enumerate(df_hab.iterrows(), start=1):
            r_idx = start_row + i - 1
            ws.cell(row=r_idx, column=1, value=i).alignment = align_center
            ws.cell(row=r_idx, column=2, value=str(row.get('nome_atirador', '')).upper()).alignment = align_left
            ws.cell(row=r_idx, column=3, value=str(row.get('cr_atirador', '')).upper()).alignment = align_center
            
            cpf_val = str(row.get('cpf_atirador', '')) if 'cpf_atirador' in row else ""
            ws.cell(row=r_idx, column=4, value=cpf_val).alignment = align_center
            
            dt_val = str(row.get('data_hora', ''))
            ws.cell(row=r_idx, column=5, value=dt_val).alignment = align_center

            for c in range(1, 6):
                cell = ws.cell(row=r_idx, column=c)
                cell.font = font_data
                cell.border = thin_border
        
        current_row = start_row + len(df_hab)
    else:
        for i in range(1, 11):
            r_idx = start_row + i - 1
            ws.cell(row=r_idx, column=1, value=i).alignment = align_center
            for c in range(1, 6):
                cell = ws.cell(row=r_idx, column=c)
                cell.font = font_data
                cell.border = thin_border
        current_row = start_row + 10

    current_row += 1
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    cell_legal = ws.cell(row=current_row, column=1)
    cell_legal.value = "A relação acima, oriunda do controle biométrico e facial desta empresa, está sendo encaminhada nos termos do art. 38, § 5º, inciso III, do Decreto nº 11.615/2023."
    cell_legal.font = font_legal
    cell_legal.alignment = align_center

    current_row += 3
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    hoje_str = obter_data_hora_atual().strftime("%d de %B de %Y")
    cell_data = ws.cell(row=current_row, column=1)
    cell_data.value = f"{CIDADE_UF_ENTIDADE}, {hoje_str}"
    cell_data.font = font_data
    cell_data.alignment = align_center

    current_row += 2
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    cell_ass_gov = ws.cell(row=current_row, column=1)
    cell_ass_gov.value = "Assinatura digital (Gov.br)"
    cell_ass_gov.font = font_subtext
    cell_ass_gov.alignment = align_center

    current_row += 1
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    cell_resp_nome = ws.cell(row=current_row, column=1)
    cell_resp_nome.value = NOME_RESPONSAVEL_ENTIDADE
    cell_resp_nome.font = font_legal
    cell_resp_nome.alignment = align_center

    current_row += 1
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    cell_resp_cargo = ws.cell(row=current_row, column=1)
    cell_resp_cargo.value = "RESPONSÁVEL PELA ENTIDADE DE TIRO"
    cell_resp_cargo.font = font_data
    cell_resp_cargo.alignment = align_center

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def gerar_pdf_relatorio_cliente(df_cliente, nome_cliente):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "RELATÓRIO DE HABITUAÇÃO DE ATIRADOR", border=0, new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.set_font("Helvetica", "", 12)
    pdf.cell(0, 8, f"Atirador: {nome_cliente.upper()}", border=0, new_x="LMARGIN", new_y="NEXT", align="C")
    
    cr_val = df_cliente['cr_atirador'].iloc[0] if 'cr_atirador' in df_cliente.columns and not df_cliente.empty else ""
    if cr_val:
        pdf.cell(0, 6, f"CR: {str(cr_val).upper()}", border=0, new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(8)
    
    for idx, row in df_cliente.iterrows():
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(0, 6, f"Registo ID: {row.get('id', '')} - Data/Hora: {row.get('data_hora', '')}", border="T", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 5, f"SIGMA: {str(row.get('sigma_atirador', 'N/A')).upper()} | Arma: {row.get('tipo_arma', '')} | Munição: {row.get('tipo_municao', '')} | Qtd: {row.get('qtd_municao', '')}", border=0, new_x="LMARGIN", new_y="NEXT")
        
        hash_val = str(row.get('hash_integridade', ''))
        pdf.set_font("Helvetica", "I", 8)
        pdf.cell(0, 5, f"Hash SHA-256: {hash_val}", border=0, new_x="LMARGIN", new_y="NEXT")
        
        ass_b64 = obter_valor_assinatura(row)
        if ass_b64:
            try:
                img_bytes = base64.b64decode(ass_b64)
                with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmp_file:
                    tmp_file.write(img_bytes)
                    tmp_path = tmp_file.name
                pdf.image(tmp_path, x=15, w=50)
                os.remove(tmp_path)
            except Exception:
                pdf.cell(0, 5, "[Assinatura Indisponível]", border=0, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(5)

    return bytes(pdf.output())

# --- CONTROLE DE ESTADO E REINICIALIZAÇÃO DO FORMULÁRIO ---
if "canvas_key" not in st.session_state:
    st.session_state["canvas_key"] = 0
if "form_version" not in st.session_state:
    st.session_state["form_version"] = 0

v = st.session_state["form_version"]

# --- CONTAINER PRINCIPAL CENTRALIZADO ---
col_esq, col_centro, col_dir = st.columns([1, 8, 1])

with col_centro:
    aba = st.radio("Selecione o Modo:", ["🎯 Registro de Habituação (Atirador)", "📊 Painel Admin / Exportar"], horizontal=True)

    if aba == "🎯 Registro de Habituação (Atirador)":
        if LOGO_B64:
            st.markdown(
                f"""
                <div style="display: flex; align-items: center; gap: 15px; margin-bottom: 10px;">
                    <img src="data:image/png;base64,{LOGO_B64}" style="height: 55px; width: auto; object-fit: contain;">
                    <h1 style="margin: 0; padding: 0; font-size: 2rem;">Registro de Habituação</h1>
                </div>
                """,
                unsafe_allow_html=True
            )
        else:
            st.title("Registro de Habituação")

        st.write("Preencha os dados da sessão de tiro e assine no campo abaixo.")

        df_hab = carregar_habituacoes_sheets()
        lista_cadastrados = obter_atiradores_existentes(df_hab)

        opcoes_rotulos = [a['rotulo'] for a in lista_cadastrados]
        opcao_selecionada = st.selectbox(
            "🔍 Buscar Atirador Cadastrado:",
            options=opcoes_rotulos,
            index=None,
            placeholder="Clique aqui e digite o Nome ou CR...",
            accept_new_options=True,
            key=f"select_atirador_{v}"
        )

        is_novo_cadastro = True
        atirador_obj = None

        if opcao_selecionada:
            match = [a for a in lista_cadastrados if a['rotulo'] == opcao_selecionada]
            if match:
                is_novo_cadastro = False
                atirador_obj = match[0]

        if is_novo_cadastro:
            nome_inicial = opcao_selecionada.upper() if opcao_selecionada else ""
            nome_input = st.text_input(
                "Nome Completo do Atirador:", 
                value=nome_inicial, 
                placeholder="Informe o nome completo para cadastrar", 
                key=f"input_nome_{v}_{opcao_selecionada}"
            )
            cr_valor = ""
            sigma_valor = ""
        else:
            nome_input = atirador_obj['nome']
            cr_valor = atirador_obj['cr']
            sigma_valor = atirador_obj['sigma']
            
            st.text_input("Nome Completo do Atirador:", value=nome_input, disabled=True, help="Atirador selecionado na busca acima.", key=f"input_nome_dis_{v}")

        col_sigma, col_cr = st.columns(2)
        with col_sigma:
            sigma_input = st.text_input(
                "Número do SIGMA:", 
                value=sigma_valor if not is_novo_cadastro else "", 
                placeholder="Informe o número do SIGMA", 
                key=f"input_sigma_{v}_{opcao_selecionada}"
            )
        with col_cr:
            cr_input = st.text_input(
                "Número do CR:", 
                value=cr_valor if not is_novo_cadastro else "", 
                placeholder="Informe o número do CR", 
                disabled=not is_novo_cadastro, 
                key=f"input_cr_{v}_{opcao_selecionada}"
            )

        col1, col2 = st.columns(2)
        with col1:
            tipo_arma = st.selectbox("Tipo de Arma:", list(OPCOES_MUNICAO_POR_ARMA.keys()), key=f"select_tipo_arma_{v}")
        with col2:
            tipo_municao = st.selectbox("Tipo/Calibre de Munição:", OPCOES_MUNICAO_POR_ARMA[tipo_arma], key=f"select_tipo_municao_{v}")

        qtd_input = st.text_input("Quantidade de Munição Utilizada:", placeholder="Informe a quantidade utilizada", key=f"input_qtd_{v}")

        st.subheader("🖋️ Assinatura Digital")
        st.caption("Assine dentro da caixa abaixo:")

        canvas_result = st_canvas(
            fill_color="rgba(255, 255, 255, 0)",
            stroke_width=3,
            stroke_color="#000000",
            background_color="#FFFFFF",
            height=200,
            drawing_mode="freedraw",
            update_streamlit=True,
            return_image_data=True,
            key=f"canvas_assinatura_{st.session_state['canvas_key']}"
        )

        if st.button("🧹 Limpar Assinatura"):
            st.session_state["canvas_key"] += 1
            st.rerun()

        st.info("📌 Registro com carimbo de tempo e hash criptográfico de validação (Lei 14.063/2020).")

        if st.button("✅ Registrar Habituação", type="primary", use_container_width=True):
            nome_final = nome_input.strip().upper() if nome_input else ""
            cr_final = cr_input.strip().upper() if cr_input else ""
            sigma_final = sigma_input.strip().upper() if sigma_input else ""

            if not nome_final:
                st.error("⚠️ Preenchimento obrigatório: Por favor, informe o Nome do Atirador.")
                st.stop()

            if not cr_final:
                st.error("⚠️ Preenchimento obrigatório: Por favor, informe o CR do Atirador.")
                st.stop()

            if not sigma_final:
                st.error("⚠️ Preenchimento obrigatório: Por favor, informe o número do SIGMA.")
                st.stop()

            if not qtd_input or not qtd_input.strip():
                st.error("⚠️ Preenchimento obrigatório: Por favor, informe a quantidade de munição utilizada.")
                st.stop()

            try:
                qtd_municao = int(qtd_input.strip())
                if qtd_municao <= 0:
                    st.error("⚠️ A quantidade de munição deve ser maior que zero.")
                    st.stop()
            except ValueError:
                st.error("⚠️ Por favor, informe um número válido para a quantidade de munição.")
                st.stop()

            if is_novo_cadastro and lista_cadastrados:
                nomes_existentes = [a['nome'].upper() for a in lista_cadastrados]
                crs_existentes = [a['cr'].upper() for a in lista_cadastrados if a['cr']]

                if nome_final in nomes_existentes:
                    st.error(f"⚠️ Atirador já cadastrado! O nome '{nome_final}' já existe no sistema. Por favor, selecione-o no campo 'Buscar Atirador Cadastrado'.")
                    st.stop()

                if cr_final in crs_existentes:
                    st.error(f"⚠️ CR já cadastrado! O CR '{cr_final}' já pertence a outro atirador. Por favor, selecione seu cadastro acima.")
                    st.stop()

            img_data = canvas_result.image_data if canvas_result is not None else None
            assinatura_valida = False
            if img_data is not None and isinstance(img_data, np.ndarray):
                if img_data.shape[2] == 4:
                    alpha = img_data[:, :, 3]
                    if np.any(alpha > 0):
                        assinatura_valida = True

            if not assinatura_valida:
                st.error("⚠️ Preenchimento obrigatório: Por favor, assine no campo de assinatura antes de salvar.")
                st.stop()

            with st.spinner("Gravando no Google Sheets..."):
                try:
                    salvar_habituação_sheets(
                        nome_final,
                        cr_final,
                        sigma_final,
                        tipo_arma,
                        tipo_municao,
                        qtd_municao,
                        img_data
                    )
                    
                    st.session_state["canvas_key"] += 1
                    st.session_state["form_version"] += 1
                    
                    st.success("HABITUALIDADE REGISTRADA COM SUCESSO!")
                    time.sleep(2)
                    st.rerun()
                except Exception as e:
                    st.error(f"Erro ao salvar: {e}")

    else:
        # --- PAINEL ADMINISTRATIVO COM FILTROS E OPÇÕES DE EXPORTAÇÃO ---
        col_titulo, col_filtro = st.columns([1.5, 1])
        with col_titulo:
            if LOGO_B64:
                st.markdown(
                    f"""
                    <div style="display: flex; align-items: center; gap: 15px; margin-bottom: 10px;">
                        <img src="data:image/png;base64,{LOGO_B64}" style="height: 50px; width: auto; object-fit: contain;">
                        <h1 style="margin: 0; padding: 0; font-size: 2rem;">Painel Administrativo</h1>
                    </div>
                    """,
                    unsafe_allow_html=True
                )
            else:
                st.title("📊 Painel Administrativo")
            
        with st.spinner("Buscando registros da nuvem..."):
            df_hab = carregar_habituacoes_sheets()

        with col_filtro:
            st.write("")
            if not df_hab.empty and 'nome_atirador' in df_hab.columns:
                lista_clientes = ["Todos os Atiradores"] + sorted([x for x in df_hab['nome_atirador'].unique() if x])
            else:
                lista_clientes = ["Todos os Atiradores"]
                
            cliente_selecionado = st.selectbox(
                "🎯 Filtrar por Atirador:", 
                options=lista_clientes,
                index=0,
                placeholder="Digite para filtrar o atirador...",
                accept_new_options=True,
                key="select_filtro_admin"
            )

            intervalo_datas = st.date_input(
                "📅 Filtrar por Período (Início - Fim):",
                value=(),
                format="DD/MM/YYYY",
                key="filtro_datas_admin"
            )

        st.write("Visualização de registros salvos no Google Sheets e exportação em Excel/PDF.")

        if df_hab.empty:
            st.info("Nenhuma habituação registrada até o momento.")
        else:
            df_filtrado = df_hab.copy()

            if cliente_selecionado and cliente_selecionado != "Todos os Atiradores":
                df_filtrado = df_filtrado[df_filtrado['nome_atirador'] == cliente_selecionado]

            str_mes_ano = "MÊS DE ________ DE 2026"

            if 'data_hora' in df_filtrado.columns and isinstance(intervalo_datas, (list, tuple)) and len(intervalo_datas) > 0:
                df_filtrado['dt_parsed'] = pd.to_datetime(df_filtrado['data_hora'], errors='coerce').dt.date
                
                if len(intervalo_datas) == 2:
                    data_ini, data_fim = intervalo_datas[0], intervalo_datas[1]
                    df_filtrado = df_filtrado[(df_filtrado['dt_parsed'] >= data_ini) & (df_filtrado['dt_parsed'] <= data_fim)]
                    str_mes_ano = f"PERÍODO DE {data_ini.strftime('%d/%m/%Y')} A {data_fim.strftime('%d/%m/%Y')}"
                elif len(intervalo_datas) == 1:
                    data_unica = intervalo_datas[0]
                    df_filtrado = df_filtrado[df_filtrado['dt_parsed'] == data_unica]
                    str_mes_ano = f"DIA {data_unica.strftime('%d/%m/%Y')}"

                df_filtrado = df_filtrado.drop(columns=['dt_parsed'], errors='ignore')

            if df_filtrado.empty:
                st.warning("Nenhum registro encontrado para os filtros selecionados.")
            else:
                cols_ocultar = [c for c in df_filtrado.columns if 'assinatura' in c.lower()]
                df_exibicao = df_filtrado.drop(columns=cols_ocultar, errors='ignore')
                st.dataframe(df_exibicao, use_container_width=True)

                st.subheader("🔍 Visualizar Registros e Assinaturas")
                
                for idx, row in df_filtrado.iterrows():
                    nome = row.get('nome_atirador', '')
                    dt = row.get('data_hora', '')
                    reg_id = row.get('id', idx)
                    
                    with st.expander(f"ID #{reg_id} - {nome} ({dt})"):
                        col_info, col_img = st.columns([1, 1])
                        
                        with col_info:
                            st.markdown(f"**Atirador:** {row.get('nome_atirador', '')}")
                            st.markdown(f"**CR:** {row.get('cr_atirador', '')}")
                            st.markdown(f"**SIGMA:** {row.get('sigma_atirador', 'N/A')}")
                            st.markdown(f"**Arma/Calibre:** {row.get('tipo_arma', '')} - {row.get('tipo_municao', '')}")
                            st.markdown(f"**Qtd. Munição:** {row.get('qtd_municao', '')}")
                            st.markdown(f"**Hash SHA-256:** `{row.get('hash_integridade', '')}`")
                        
                        with col_img:
                            st.markdown("**Assinatura Capturada:**")
                            ass_b64 = obter_valor_assinatura(row)
                            if ass_b64:
                                try:
                                    img_bytes = base64.b64decode(ass_b64)
                                    st.image(img_bytes, width=280)
                                except Exception:
                                    st.error("Erro ao converter imagem Base64.")
                            else:
                                st.caption("Sem imagem de assinatura válida nesta linha.")

                st.markdown("---")
                
                # BOTÕES DE EXPORTAÇÃO
                st.subheader("📥 Exportação de Relatórios")
                col_btn1, col_btn2, col_btn3 = st.columns([1, 1, 1])

                with col_btn1:
                    excel_sfpc = gerar_excel_modelo_sfpc(df_filtrado, periodo_mes_ano=str_mes_ano)
                    st.download_button(
                        label="📋 Relação Oficial SFPC (Excel)",
                        data=excel_sfpc,
                        file_name=f"relacao_oficial_sfpc_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary",
                        use_container_width=True
                    )

                with col_btn2:
                    excel_data = gerar_excel_com_assinaturas(df_filtrado)
                    st.download_button(
                        label="📥 Planilha Completa c/ Assinaturas",
                        data=excel_data,
                        file_name=f"habituacoes_com_assinaturas_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True
                    )

                with col_btn3:
                    if cliente_selecionado and cliente_selecionado != "Todos os Atiradores" and not df_filtrado.empty:
                        pdf_data = gerar_pdf_relatorio_cliente(df_filtrado, cliente_selecionado)
                        st.download_button(
                            label=f"📄 Relatório PDF ({cliente_selecionado})",
                            data=pdf_data,
                            file_name=f"relatorio_habituação_{cliente_selecionado.replace(' ', '_')}.pdf",
                            mime="application/pdf",
                            use_container_width=True
                        )
                    else:
                        st.button(
                            "📄 Selecione um Atirador para PDF",
                            disabled=True,
                            use_container_width=True,
                            help="Selecione um atirador específico no filtro do topo para gerar o relatório em PDF."
                        )

    # --- RODAPÉ INSTITUCIONAL DE AUTORIA ---
    st.markdown("---")
    st.markdown(
        "<div style='text-align: center; color: #777777; font-size: 13px; margin-top: 10px;'>"
        "💻 <b>Desenvolvido por RODRIGO HENRIQUE NEVES</b> &copy; 2026 | Sistema de Habituação de Atiradores"
        "</div>",
        unsafe_allow_html=True
    )
